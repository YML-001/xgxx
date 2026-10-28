#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文风格 Word 排版工具箱（GB/T 9704 惯例）。

供 tools/ 下各 docx 生成脚本共用：字体常量、单元格底色、表格工具与 Builder。
四份文档（PRD / 功能介绍 / 功能清单与工作量评估 / 消息提醒与模板）统一复用本模块，
保证版式一致；修改排版只需改此处一份。
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_WEST = "Times New Roman"
FONT_HEAD = "黑体"
FONT_BODY = "宋体"
FONT_KAI = "楷体"

CLR_HEAD_BG = "DCE9F5"      # 表头底色（中南蓝浅色）
CLR_SUB_BG = "F2F7FB"       # 小计行底色
CLR_ACCENT = RGBColor(0x1F, 0x5C, 0x99)



# ---------------------------------------------------------------- 基础排版工具

def set_run(run, cn_font=FONT_BODY, size=12, bold=False, color=None):
    run.font.name = FONT_WEST
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    if color is not None:
        run.font.color.rgb = color
    return run


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def set_widths(table, widths_cm):
    table.autofit = False
    fixed_layout(table)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def cell_text(cell, text, size=9.5, bold=False, align="center",
              cn_font=FONT_BODY, color=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run(p.add_run(str(text)), cn_font, size, bold, color)


class Builder:
    def __init__(self):
        self.doc = Document()
        self._init_styles()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(s, attr, Cm(2.5))

    def _init_styles(self):
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT_WEST
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12.5)):
            st = self.doc.styles[name]
            st.font.name = FONT_WEST
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_HEAD)
            st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    # -------------------------------------------------------------- 内容元素
    def blank(self, n=1):
        for _ in range(n):
            self.doc.add_paragraph()

    def h(self, text, level=1):
        return self.doc.add_paragraph(text, style=f"Heading {level}")

    def body(self, text, indent=True, size=12, cn_font=FONT_BODY, align="justify"):
        p = self.doc.add_paragraph()
        p.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if align == "justify"
                       else getattr(WD_ALIGN_PARAGRAPH, align.upper()))
        if indent:
            p.paragraph_format.first_line_indent = Pt(size * 2)
        set_run(p.add_run(text), cn_font, size)
        return p

    def bullet(self, text, size=12):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(text), FONT_BODY, size)
        return p

    def caption(self, text, size=10.5):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(text), FONT_HEAD, size, bold=True)
        return p

    def note(self, text, size=10.5):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(text), FONT_KAI, size, color=RGBColor(0x55, 0x55, 0x55))
        return p

    def table(self, headers, rows, widths, font=9.5, aligns=None,
              highlight_rows=(), head_font=None):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(headers):
            c = t.rows[0].cells[i]
            cell_text(c, htxt, size=head_font or font, bold=True,
                      cn_font=FONT_HEAD, color=CLR_ACCENT)
            shade(c, CLR_HEAD_BG)
        repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            emph = ri in highlight_rows
            for i, val in enumerate(row):
                a = (aligns[i] if aligns else "center")
                cell_text(cells[i], val, size=font, bold=emph, align=a)
                if emph:
                    shade(cells[i], CLR_SUB_BG)
        set_widths(t, widths)
        return t

    # -------------------------------------------------------------- 节与页眉脚
    def new_section(self, landscape=False, margin=2.5):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Cm(29.7), Cm(21)
        else:
            s.orientation = WD_ORIENT.PORTRAIT
            s.page_width, s.page_height = Cm(21), Cm(29.7)
        for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(s, attr, Cm(margin))
        s.footer.is_linked_to_previous = False
        return s

    def page_number_footers(self):
        """页脚居中显示「第 X 页」；PAGE 为域，由 Word 渲染时自动计算。"""
        for s in self.doc.sections[1:]:
            p = s.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            set_run(p.add_run("第 "), FONT_BODY, 9)

            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), " PAGE \\* MERGEFORMAT ")
            inner = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), FONT_WEST)
            fonts.set(qn("w:hAnsi"), FONT_WEST)
            fonts.set(qn("w:eastAsia"), FONT_BODY)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "18")  # 半磅值，18 = 9pt
            rpr.append(fonts)
            rpr.append(sz)
            t = OxmlElement("w:t")
            t.text = "1"
            inner.append(rpr)
            inner.append(t)
            fld.append(inner)
            p._p.append(fld)

            set_run(p.add_run(" 页"), FONT_BODY, 9)

    def toc(self):
        p = self.doc.add_paragraph()
        run = p.add_run()
        set_run(run, FONT_BODY, 12)
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = r'TOC \o "1-2" \h \z \u'
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "（请在 Word 中右键此处选择「更新域」生成目录）"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for el in (b, instr, sep, placeholder, end):
            run._r.append(el)

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
        print(f"已生成: {path}")
