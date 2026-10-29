#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《学工信箱建设方案》，供领导确认后作为开发依据。"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_kit import Builder, set_run, FONT_BODY, FONT_HEAD

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "学工信箱建设方案.docx")
FLOW = os.path.join(ROOT, "docs", "manual-assets", "mailbox-flow.png")

def lines(b, text):
    for line in text.split("\n"):
        b.body(line, indent=False)


def figure(b, path, caption):
    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(16))
    cap = b.doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(cap.add_run(caption), FONT_HEAD, 10.5, bold=True)


def _lvl(ilvl, start, num_fmt, text, style_id, legal=False):
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), str(ilvl))
    for tag, val in (("w:start", str(start)), ("w:numFmt", num_fmt)):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    if legal:
        lvl.append(OxmlElement("w:isLgl"))
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), style_id)
    lvl.append(ps)
    lt = OxmlElement("w:lvlText")
    lt.set(qn("w:val"), text)
    lvl.append(lt)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "space")
    lvl.append(suff)
    return lvl


def setup_heading_numbers(b):
    """一级用「一、」，二级用「1.1」，由 Word 编号生成，目录可随标题更新。"""
    root = b.doc.part.numbering_part._element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "10")
    ml = OxmlElement("w:multiLevelType")
    ml.set(qn("w:val"), "multilevel")
    abstract.append(ml)
    abstract.append(_lvl(0, 1, "chineseCounting", "%1、", "Heading1"))
    abstract.append(_lvl(1, 1, "decimal", "%1.%2", "Heading2", legal=True))
    abstract.append(_lvl(2, 1, "decimal", "%1.%2.%3", "Heading3", legal=True))
    root.insert(0, abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "10")
    aid = OxmlElement("w:abstractNumId")
    aid.set(qn("w:val"), "10")
    num.append(aid)
    root.append(num)
    for style_name, ilvl in (("Heading 1", 0), ("Heading 2", 1), ("Heading 3", 2)):
        pPr = b.doc.styles[style_name].element.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(ilvl))
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), "10")
        numPr.append(il)
        numPr.append(nid)
        pPr.append(numPr)


def toc(b):
    title = b.doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    set_run(title.add_run("目  录"), FONT_HEAD, 16, bold=True)

    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        return r

    def fld(kind):
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
        return run_with(node)

    p = b.doc.add_paragraph()
    p._p.append(fld("begin"))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    p._p.append(run_with(instr))
    p._p.append(fld("separate"))
    shown = OxmlElement("w:t")
    shown.text = "打开文档后请更新目录。目录按标题自动生成序号和页码。"
    p._p.append(run_with(shown))
    p._p.append(fld("end"))
    b.note("本目录是 Word 目录域。用 Microsoft Word 或 WPS 打开时，若提示更新域，请选择「更新整个目录」。章节序号由标题样式自动编排，增删章节后再次更新即可。")
    b.doc.add_page_break()


def enable_update(b):
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    b.doc.settings.element.append(el)


def fn_table(b, rows):
    b.table(
        ["编号", "功能", "说明", "使用人"],
        rows,
        [1.6, 3.4, 7.6, 3.4],
        font=10,
        aligns=["center", "center", "left", "center"],
    )



def main():
    from mailbox_terms import assert_catalog, assert_docx, TEMPLATES
    from proposal_fill import fill

    assert_catalog()
    b = Builder()
    enable_update(b)
    setup_heading_numbers(b)
    title = b.doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    set_run(title.add_run("学工信箱建设方案"), FONT_HEAD, 22, bold=True)
    b.note("中南大学智慧学工　·　2026年10月　·　供领导确认。确认后按本方案建设。")
    b.body("学工信箱是学生向学生工作部反映建议、问题、咨询和投诉的统一入口。来信先由综合办公室受理并指定负责科室，分派给科室负责人办理。正式答复经科室分管领导审批后，由科室负责人确认回复学生。本文说明本板块的功能和办理流程，作为确认稿。")
    toc(b)
    fill(b, figure, lines, fn_table, FLOW)
    b.save(OUT)
    doc = assert_docx(OUT, expect_toc=True)
    print("段落数: %d" % len(doc.paragraphs))
    print("表格数: %d" % len(doc.tables))
    print("模板条数: %d" % len(TEMPLATES))
    print("校验通过", OUT)


if __name__ == "__main__":
    main()
