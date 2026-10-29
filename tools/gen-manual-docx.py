#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按《视频库-学习端操作手册》的版式生成学工信箱操作手册。

输出四份：
  学工信箱-操作手册（移动端+PC端）.docx
  学工信箱-学生端操作手册（移动端+PC端）.docx
  学工信箱-科室教师端操作手册（移动端+PC端）.docx
  学工信箱-综合办管理员操作手册（移动端+PC端）.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "docs", "manual-assets", "prod")
OUT_DIR = os.path.join(ROOT, "docs")

FONT_WEST = "Times New Roman"
FONT_TITLE = "微软雅黑"
FONT_BODY = "宋体"


def set_run(run, cn, size, bold=False):
    run.font.name = FONT_WEST
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    return run


class Manual:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for name in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, name, Cm(2.5))
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT_WEST
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = pf.space_after = Pt(0)
        for name, size in (("Heading 1", 15), ("Heading 2", 13)):
            st = self.doc.styles[name]
            st.font.name = FONT_WEST
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_TITLE)
            st.paragraph_format.space_before = Pt(6)
            st.paragraph_format.space_after = Pt(3)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        settings = self.doc.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

    def cover(self, subtitle):
        for _ in range(6):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run("中南大学智慧学工 · 学工信箱"), FONT_TITLE, 22, True)
        p2 = self.doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p2.add_run(subtitle), FONT_TITLE, 18, True)
        for _ in range(8):
            self.doc.add_paragraph()
        p3 = self.doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p3.add_run("2026 年 10 月"), FONT_BODY, 12)
        self.doc.add_page_break()

    def h1(self, text):
        p = self.doc.add_paragraph(style="Heading 1")
        set_run(p.add_run(text), FONT_TITLE, 15, True)
        return p

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        set_run(p.add_run(text), FONT_TITLE, 13, True)
        return p

    def toc(self, lines):
        """插入 Word 目录域（TOC），打开时由 Office 更新页码。lines 为打开前可见的目录缓存。"""
        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(8)
        set_run(title.add_run("目  录"), FONT_TITLE, 16, True)

        p = self.doc.add_paragraph()
        run = p.add_run()
        set_run(run, FONT_BODY, 12)

        def fld(el_type):
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), el_type)
            r = OxmlElement("w:r")
            r.append(node)
            return r

        begin = fld("begin")
        instr_r = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = r' TOC \o "1-2" \h \z \u '
        instr_r.append(instr)
        sep = fld("separate")
        p._p.append(begin)
        p._p.append(instr_r)
        p._p.append(sep)

        for level, text in lines:
            item = self.doc.add_paragraph()
            item.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            item.paragraph_format.space_before = Pt(2)
            item.paragraph_format.space_after = Pt(2)
            if level == 2:
                item.paragraph_format.left_indent = Cm(0.75)
            set_run(item.add_run(text), FONT_BODY, 12, level == 1)

        end_p = self.doc.add_paragraph()
        end_p._p.append(fld("end"))
        note = self.doc.add_paragraph()
        note.paragraph_format.space_before = Pt(10)
        set_run(note.add_run(
            "本目录为 Word 目录域。使用 Microsoft Word 或 WPS 打开时，若提示更新域，请选择“更新”，页码会按正文标题自动生成。"
        ), "楷体", 10.5)
        self.doc.add_page_break()

    def page_numbers(self):
        for sec in self.doc.sections:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            set_run(p.add_run("第 "), FONT_BODY, 9)
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), " PAGE ")
            inner = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = "1"
            inner.append(t)
            fld.append(inner)
            p._p.append(fld)
            set_run(p.add_run(" 页"), FONT_BODY, 9)

    def p(self, text):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Pt(24)
        set_run(para.add_run(text), FONT_BODY, 12)

    def image(self, name, caption, width_cm):
        path = os.path.join(ASSETS, name)
        if not os.path.exists(path):
            raise FileNotFoundError(path)
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.add_run().add_picture(path, width=Cm(width_cm))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(6)
        set_run(cap.add_run(caption), FONT_TITLE, 11)

    def table(self, headers, rows, widths):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(cell.paragraphs[0].add_run(h), FONT_TITLE, 11, True)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_run(cells[i].paragraphs[0].add_run(val), FONT_BODY, 11)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    def save(self, filename):
        path = os.path.join(OUT_DIR, filename)
        self.doc.save(path)
        print("已生成:", path)


def fig_width(name):
    if name.startswith("m-") or name.startswith("portal-wx"):
        return 6.6
    return 15.0


def emit_steps(doc, chapter, items, start=1):
    """每个步骤：二级标题 + 说明 + 一张图。图号与小节号一致。"""
    n = start
    for it in items:
        doc.h2(f"{chapter}.{n} {it['title']}")
        doc.p(it["body"])
        doc.image(it["img"], f"图 {chapter}-{n}  {it['cap']}", fig_width(it["img"]))
        n += 1
    return n


# ---------------------------------------------------------------- 共用素材

RULES_ALL = [
    ["统一受理", "学生提交后，来信一律送达综合办公室，状态为「待分派」，此时尚未指定承办科室与承办人。"],
    ["派发到人", "综合办先选择承办科室，再选择该科室处于启用状态的成员。没有启用成员的科室不能选。"],
    ["建议归口", "系统按事务分类预选建议科室，并把负责该事务的成员标为「对口」。实际派给谁由综合办确定，可以改派。"],
    ["办理权限", "科室教师可以查看本科室已经派发的全部信件，只能对派发给本人的信件填写答复并办结。"],
    ["待分派范围", "尚未派发的信件只在综合办可见。科室教师看不到，也不能办理。"],
    ["办结时限", "时限从来信提交之日起算。超过配置天数（默认 7 天）仍未派发或未办结，自动标记「超期」。"],
    ["催办送达", "已经派发的催办，送达承办科室和承办人；尚未派发的，送达综合办公室。"],
    ["评价与二次办理", "办结后学生先做满意度评价。不满意可申请二次办理，信件回到「办理中」，仍由原承办人继续办理。"],
    ["一事一信", "标题相同的已提交信件不能重复提交。正文不少于 10 字、不超过 2000 字。"],
]

RULES_STU = [r for r in RULES_ALL if r[0] in ("统一受理", "办结时限", "催办送达", "评价与二次办理", "一事一信")]
RULES_SEC = [r for r in RULES_ALL if r[0] in ("办理权限", "待分派范围", "办结时限", "催办送达", "评价与二次办理")]
RULES_ADM = [r for r in RULES_ALL if r[0] in ("统一受理", "派发到人", "建议归口", "办理权限", "办结时限", "催办送达")]

ROLES = [
    ["学生", "本人信件，以及公开信件、制度文件", "写信、在我的信件查看进度、撤回、补充材料、催办、评价、申请二次办理"],
    ["科室教师", "本科室已经派发的信件", "答复并办结派给本人的信件，查看同事承办情况，必要时重新打开"],
    ["综合办管理员", "全部信件，包括尚未派发的", "受理派发、转派、督办催办、来信选登、统计分析与后台配置"],
]

TERMINAL = [
    ["学生来信、我的信件、公开信件、催办、评价", "可办理", "可办理"],
    ["科室教师答复并办结", "可办理", "可办理"],
    ["单封派发、转派、督办提醒", "可办理", "可办理"],
    ["批量派发、批量催办", "可办理", "请在 PC 端操作"],
    ["来信选登、完整统计分析", "可办理", "移动端提供轻量统计入口"],
    ["分类、科室成员、制度、催办模板、常见问题", "可办理", "请在 PC 端配置"],
]

ENTRY_LOGIN = (
    "智慧学工系统采用学校统一身份认证，没有单独的学工信箱账号。"
    "使用学号或工号及统一身份认证密码登录即可。首次使用请先在统一身份认证平台确认账号已经激活。"
    "登录后按账号身份进入对应功能：学生用来信、查看我的信件和公开信件，科室承办人办理派给本人的信件，综合办管理员负责受理派发、督办和配置。"
    "超级管理员进入后，可以看到学生、科室和综合办的全部菜单。"
)

ENTRY_PC = (
    "PC 端通过浏览器访问智慧学工系统门户，登录后从顶部导航进入对应业务模块。"
    "建议使用 Chrome、Edge 等主流浏览器的较新版本，校内、校外网络均可直接访问。"
    "可从中南大学官网“学生天地”栏目中的“智慧学工系统”入口进入。"
)

ENTRY_PC_NEXT = (
    "点击入口后进入统一身份认证登录，随后在“选择系统”中进入“智慧学工”。"
    "进入工作台后，从顶部“业务助手”搜索“学工信箱”或“信箱”，进入本模块。"
)

ENTRY_MOBILE = (
    "移动端通过“中南大学智慧学工”小程序使用，支持 iOS 12.0、Android 8.0 及以上系统。"
    "在微信中搜索“中南大学智慧学工”，进入小程序。"
)

ENTRY_MOBILE_NEXT = (
    "进入小程序后，点击底部“服务大厅”，找到“学工信箱”并进入。"
    "需要本人处理的派发、答复、催办，也会出现在小程序“待办事项”中。"
)

# ---------------------------------------------------------------- PC 步骤

PC_STU = [
    dict(title="工作台", cap="PC 端·工作台（学生）", img="pc-stu-home.png",
         body="进入学工信箱后打开工作台，可查看我的来信、我的已办结和我的满意度，并从快捷入口进入写信、我的信件和公开信件。页面下方展示最新公开信件和本人信件的办理情况。"),
    dict(title="填写并提交信件", cap="PC 端·我要写信", img="pc-stu-write.png",
         body="在左侧点击“我要写信”。填写姓名、证件信息、联系电话、邮箱、所在单位、联系地址、来信标题、信件分类、事务分类和正文。正文不超过 2000 字。附件支持常见图片、办公文档和压缩包，单个不超过 20MB，最多 5 个。确认后点击“提交信件”，也可以“暂存草稿”或“取消”。信件是否对外公开，由综合办在“来信选登”中决定。同一事项办结后，请勿以相同事实重复提交。"),
    dict(title="我的信件", cap="PC 端·我的信件", img="pc-stu-my.png",
         body="“我的信件”可按信件标题、信件内容和答复科室查询，并用“全部 / 待分派 / 办理中 / 已办结 / 草稿”切换。列表显示承办科室和承办人。已办结信件可在操作列申请“二次办理”；已撤回的信件仍可查看。点击信件编号打开详情。"),
    dict(title="查看信件详情", cap="PC 端·信件详情", img="pc-stu-detail.png",
         body="详情中查看基本信息、信件内容和办理流程。流程包括提交来信、归口、答复办结和满意度评价。对已办结结果不满意时，点击“对结果不满意，申请二次办理”。查看完毕点击“关闭”。"),
    dict(title="公开信件", cap="PC 端·公开信件", img="pc-stu-public.png",
         body="“公开信件”展示综合办已经选登的来信和答复，可按标题、内容和答复科室筛选。未经选登的信件不会出现在这里。"),
    dict(title="制度文件", cap="PC 端·制度文件", img="pc-stu-policy.png",
         body="在“制度文件”中按类别浏览现行制度，点击文件查看正文和附件。文件由综合办在后台维护。"),
    dict(title="数据统计", cap="PC 端·数据统计", img="pc-stu-stats.png",
         body="“数据统计”展示我的来信、已办结、平均办理时长和满意度，以及分类分布、承办科室来信量。页面不展示他人信件正文。"),
    dict(title="科室联系方式", cap="PC 端·科室联系方式", img="pc-stu-dept.png",
         body="“科室联系方式”列出各科室负责人、职责和承接事务，便于了解分工和线下联系。来信仍由综合办受理后再派发到承办人，学生不在写信时指定承办老师。"),
    dict(title="常见问题", cap="PC 端·常见问题", img="pc-stu-faq.png",
         body="“常见问题”按办理流程、使用指南、受理范围展示管理员维护的问答。没有找到答案时，可从页面进入写信。"),
]

PC_SEC = [
    dict(title="科室工作台", cap="PC 端·科室工作台", img="pc-sec-home.png",
         body="在左侧“本科室办理”中打开“科室工作台”。页面按本人承办统计累计、办理中、已办结和超期，并列出待办。下方是本科室成员的承办分布，同事的信件可以查看。若账号同时担任其他职务，左侧还会出现办理或配置菜单，那些操作见综合办手册。"),
    dict(title="科室收件箱", cap="PC 端·科室收件箱", img="pc-sec-inbox.png",
         body="“科室收件箱”可在“我承办的”和“本科室全部”之间切换，并按全部、办理中、已办结筛选。可按标题、分类和事务查询，也可导出列表。办理中的信件点击“办理”；已办结的信件点击“详情”。"),
    dict(title="查看并办理信件", cap="PC 端·信件详情（科室）", img="pc-sec-reply.png",
         body="打开信件后可查看来信内容、附件、来信人信息、催办记录和派发记录。已办结信件以查看为主。办理中的信件在列表中点击“办理”，填写答复后再办结；答复不能为空。需要更换承办人时，由综合办转派。"),
]

PC_ADM = [
    dict(title="办理工作台", cap="PC 端·办理工作台", img="pc-adm-home.png",
         body="综合办管理员登录后进入办理工作台，查看待分派、办理中、已办结和超期数量。待分派列表可直接派发。全部来信都先集中在综合办公室，再由本岗位派到具体承办老师。"),
    dict(title="待分派信件", cap="PC 端·待分派信件", img="pc-adm-dispatch.png",
         body="“待分派信件”集中展示尚未指定承办人的来信，并统计待分派总数、其中超期件数和可选承办老师数。可按标题、事务筛选，列表给出建议归口。页面下方保留最近派发记录。"),
    dict(title="派发到承办老师", cap="PC 端·派发信件", img="pc-adm-dispatch-modal.png",
         body="点击“派发”打开“派发承办人”窗口。先选择承办科室，建议归口会预先标出，并显示各科室在办数量；再选择承办人，可填写派发说明（不超过 200 字）。核对无误后点击“确认派发”，未确定时点击“取消”。列表中勾选多封后，可使用“批量派发”，统一派给同一承办人。"),
    dict(title="信件收件箱", cap="PC 端·信件收件箱（综合办）", img="pc-adm-inbox.png",
         body="“信件收件箱”查看全部信件，页签包含待分派，可按承办人检索。待分派的行可以直接派发；已经派发、尚未办结的信件可以“转派他人”。转派同样先选科室再选成员，原承办信息保留在派发记录中，办结时限按转派时间重算。"),
    dict(title="督办催办", cap="PC 端·督办催办", img="pc-adm-supervise.png",
         body="“督办催办”汇总超期未办结、超期未派发、学生催办和临期待办。列表给出承办科室、办结时限和状态。对待分派信件可继续派发，对办理中信件可进入办理。勾选多封后可“批量催办”。"),
    dict(title="来信选登管理", cap="PC 端·来信选登管理", img="pc-adm-publish.png",
         body="在“来信选登管理”中查看已办结信件。状态为“未公开”时，点击“选登公开”后，学生端“公开信件”可以看到该信和答复。点击行可先查看详情。是否公开由综合办选登决定，学生写信时不单独勾选。"),
    dict(title="统计分析", cap="PC 端·统计分析", img="pc-adm-analytics.png",
         body="“统计分析”按年、季、月查看来信趋势，并可按信件分类、承办科室筛选。页面同时展示待分派数量、办结率、好评率、分类占比、科室承办量和承办人绩效，用于调配派发和督办重点。"),
    dict(title="分类管理", cap="PC 端·分类管理", img="pc-adm-category.png",
         body="“分类管理”维护信件分类和事务分类，并把事务分类绑定到建议归口科室。该绑定只用于派发时预选科室和标记对口成员，不代替人工派发。"),
    dict(title="科室与成员管理", cap="PC 端·科室管理", img="pc-adm-section.png",
         body="“科室管理”维护科室负责人、电话、邮箱和职责。点击科室后，在下方维护承办老师：姓名、岗位、电话、邮箱、负责事务，以及启用或停用。只有启用中的成员会出现在派发窗口。仍有办理中信件的成员，需先转派其在办件，才能停用；已有承办记录的成员可以停用，不能删除。综合办公室作为统一受理科室，不能删除。"),
    dict(title="制度文件管理", cap="PC 端·制度文件管理", img="pc-adm-policy.png",
         body="“制度文件管理”新增、修改或删除制度文件，维护文号、标题、类别、效力和正文，并可上传附件。保存后，学生端和科室端的“制度文件”同步更新。"),
    dict(title="超期与催办配置", cap="PC 端·超期与催办配置", img="pc-adm-sla.png",
         body="在“超期与催办配置”中设置超过多少天未办结视为超期，以及催办消息模板和默认渠道。模板可使用科室、承办人、姓名、事务分类、标题、编号、来信时间、已过天数、催办时间等占位符。新来信、派发、办结、催办等提醒会发送站内消息，并同时发送 ITS；此处勾选的渠道是在此基础上追加。修改后点击“保存配置”。将天数改小后，更多未办结信件会标为超期。"),
    dict(title="常见问题配置", cap="PC 端·常见问题配置", img="pc-adm-faq.png",
         body="“常见问题配置”维护问答的问题、解答和分类，可启用、停用，并上移、下移调整顺序。学生端只展示启用项，并按这里的顺序排列。"),
]

MO_STU = [
    dict(title="首页", cap="移动端·首页（学生）", img="m-stu-home.png",
         body="首页展示我的来信、我的已办结和我的满意度，以及写信、我的、公开三个快捷功能，还有最新公开信件和办理占比。底部为“首页 / 公开 / 我的 / 更多”。"),
    dict(title="写信须知", cap="移动端·写信须知", img="m-stu-notice.png",
         body="点击“写信”后先阅读须知。须知说明一事一信、受理范围和不予受理事项，并列出“提交来信—综合办受理—派发承办—科室办理—答复反馈”。倒计时结束后再填写。信件是否公开由综合办选登决定。填写项和提交方式与 PC 端相同，可暂存草稿或提交。"),
    dict(title="我的信件", cap="移动端·我的信件", img="m-stu-my.png",
         body="“我的”按全部、待分派、办理中、已办结、草稿查看。卡片显示分类、承办科室和日期。已办结信件可申请二次办理，也可查看详情。"),
    dict(title="公开信件", cap="移动端·公开信件", img="m-stu-public.png",
         body="底部“公开”浏览综合办已选登的来信和答复。"),
    dict(title="更多功能", cap="移动端·更多", img="m-stu-more.png",
         body="“更多”进入制度文件、数据统计、科室联系和常见问题，内容与 PC 端对应菜单一致。"),
]

MO_SEC = [
    dict(title="科室工作台", cap="移动端·科室工作台", img="m-sec-home.png",
         body="进入学工信箱后打开科室工作台，显示本人承办累计、办理中、已办结和超期，以及本科室成员承办分布。底部为“工作台 / 收件箱 / 我的”。"),
    dict(title="科室收件箱", cap="移动端·科室收件箱", img="m-sec-inbox.png",
         body="收件箱可切换“我承办的”和“本科室全部”，并按全部、办理中、已办结筛选。点击“查看详情”阅读来信、催办和派发记录。办理中的信件在此办理并答复。"),
]

MO_ADM = [
    dict(title="办理工作台", cap="移动端·办理工作台", img="m-adm-home.png",
         body="进入学工信箱后打开综合办工作台，显示待分派、办理中、已办结和超期，并列出待分派来信和建议归口。点击“派发”即可指定承办科室和承办人，确认前可以取消。批量派发、批量催办和后台配置请使用 PC 端。"),
]


def entry_chapter(doc, who):
    doc.h1("二、系统入口与登录")
    doc.h2("2.1 登录方式")
    doc.p(ENTRY_LOGIN)
    doc.h2("2.2 PC 端入口")
    doc.p(ENTRY_PC)
    doc.image("portal-pc.png", "图 2-1  中南大学官网“学生天地”→“智慧学工系统”入口", 15)
    doc.p(ENTRY_PC_NEXT)
    doc.h2("2.3 移动端入口")
    doc.p(ENTRY_MOBILE)
    doc.image("portal-wx.jpeg", "图 2-2  微信搜索“中南大学智慧学工”小程序", 6.6)
    doc.p(ENTRY_MOBILE_NEXT)
    doc.h2("2.4 终端适用说明")
    if who == "stu":
        doc.p("学生来信、我的信件、公开信件、催办、补充、评价和二次办理，在 PC 端与移动端都可以办理，进度按同一账号互通。移动端写信前会先阅读写信须知。")
    elif who == "sec":
        doc.p("科室承办人查看收件箱、办理答复，在 PC 端与移动端都可以进行。公开信件和制度文件两端均可查阅。批量操作不在移动端提供。")
    elif who == "adm":
        doc.p("单封派发、转派和督办提醒在 PC 端与移动端都可以办理。批量派发、批量催办、来信选登、完整统计分析，以及分类、科室成员、制度、催办模板、常见问题等配置，请在 PC 端操作。")
    else:
        doc.p("学生来信与跟踪、科室答复、单封派发和督办，PC 端与移动端都可以办理，数据按账号互通。批量操作和后台配置集中在 PC 端。")
        doc.table(["功能", "PC 端", "移动端"], TERMINAL, [7.2, 4.2, 4.6])


def overview(doc, who):
    doc.h1("一、系统概述")
    doc.h2("1.1 系统简介")
    if who == "stu":
        doc.p("学工信箱供学生向学生工作部反映建议、问题、咨询、感谢和投诉。来信统一送达综合办公室，再派发到具体科室的承办老师办理并答复。学生可随时查看进度、补充材料、催办，办结后进行评价。PC 端与移动端功能一致，数据按账号互通。")
    elif who == "sec":
        doc.p("科室教师在学工信箱中办理综合办派发给本人的学生来信：核实情况、填写答复并办结。可以查看本科室其他成员的承办信件，便于协同，但不能代为答复。PC 端与移动端均可办理，数据按账号互通。")
    elif who == "adm":
        doc.p("综合办管理员统一受理全部学生来信，研判后派发到具体科室的承办老师，并负责转派、督办、选登和统计分析。来信在派发之前处于“待分派”，只有本岗位可以处理和查看。批量操作与后台配置在 PC 端完成，移动端可处理单封派发和督办。")
    else:
        doc.p("学工信箱供学生向学生工作部提交建议、问题、咨询、感谢和投诉。来信统一送达综合办公室受理，由综合办管理员派发到具体科室的承办老师办理并答复，形成“提交—受理—派发—办理—答复—评价”的闭环。系统提供 PC 端与移动端，同一账号数据互通。本文覆盖学生、科室承办人、综合办管理员；超级管理员登录后可使用上述全部菜单。另有三份分角色手册可单独发放。")
    doc.h2("1.2 使用角色" if who == "all" else "1.2 本角色要做的事")
    if who == "all":
        doc.p("三个角色的可见范围和职责如下。账号登录后进入对应菜单，不需要在系统内切换身份。")
        doc.table(["角色", "可见范围", "主要操作"], ROLES, [3.2, 5.6, 7.2])
    elif who == "stu":
        doc.p("学生可以查看和办理本人提交的信件，并浏览公开信件、制度文件、办理统计、科室联系方式和常见问题。不能查看其他同学的未公开信件，也不能指定承办老师。")
    elif who == "sec":
        doc.p("科室教师可以查看本科室已经派发的信件，只能答复派发给本人的信件。待分派信件、其他科室的信件，以及派发、转派、后台配置，不在本角色的操作范围内。")
    else:
        doc.p("综合办管理员可以查看和办理全部信件，负责派发到人、转派、督办催办、来信选登，以及分类、科室成员、制度文件、超期催办和常见问题的配置。答复正文由承办老师填写。")
    doc.h2("1.3 核心办理规则（重要）")
    rules = {"all": RULES_ALL, "stu": RULES_STU, "sec": RULES_SEC, "adm": RULES_ADM}[who]
    doc.table(["规则", "说明"], rules, [3.6, 12.4])


FAQ_ALL = [
    ("提交之后为什么还是“待分派”？",
     "来信先送到综合办公室登记。综合办选定承办科室和承办老师并派发后，状态才变为“办理中”。派发之前，科室教师看不到这封信。"),
    ("为什么不能替同事答复？",
     "只有派发给本人的信件才能填写答复并办结。需要更换承办人时，由综合办在收件箱中转派，转派记录会保留在信件详情里。"),
    ("催办会发给谁？",
     "已经派发的信件，催办送达承办科室和承办人；还在待分派的，催办送达综合办公室。发送后该信标记为“已催办”。"),
    ("对办理结果不满意怎么办？",
     "先完成满意度评价，星级为必填。评价之后可以申请二次办理，信件回到“办理中”，仍由原承办人继续办理。如需换人，由综合办转派。"),
    ("什么时候不能撤回？",
     "来信人可以撤回“待分派”和“办理中”的本人信件。已经办结的信件不能撤回，可以评价或申请二次办理。撤回后仍可在“我的信件”中查看。"),
    ("手机上找不到批量派发和配置菜单？",
     "批量派发、批量催办、来信选登，以及分类、科室成员、制度、催办模板、常见问题的维护，在 PC 端操作。手机端可以完成单封派发、转派、督办和答复办结。"),
    ("换一台电脑或手机，原来的信件还在吗？",
     "在。信件按登录账号保存，PC 端与移动端互通，用同一账号即可继续查看和办理。"),
    ("哪些事项不予受理？",
     "非学生工作部职责范围的事项、紧急救助、涉及国家秘密或个人隐私、应当走法定程序的事项，以及重复提交的无理诉求，不予受理。具体以写信须知为准。"),
]

FAQ_STU = [FAQ_ALL[i] for i in (0, 2, 3, 4, 6, 7)]
FAQ_SEC = [
    FAQ_ALL[1],
    FAQ_ALL[2],
    ("学生申请二次办理之后，信件到哪里了？",
     "二次办理后信件回到“办理中”，仍由原承办人继续办理。请在收件箱的“办理中”中打开并重新答复。"),
    ("已办结的信件在哪里看？",
     "在科室收件箱切换到“已办结”，点击“详情”或“查看详情”，可查看来信、附件、催办记录和派发记录。"),
    FAQ_ALL[6],
]
FAQ_ADM = [
    FAQ_ALL[0],
    ("派发时怎样参考建议科室？",
     "派发窗口会标出该事务的建议归口科室，并显示各科室、各承办人的在办数量。建议只供参考，可以改派到其他科室或其他成员。确认前可以取消。"),
    ("成员无法停用或无法删除？",
     "该成员仍有办理中的信件时，需先把在办件转派给他人，才能停用。已经产生承办记录的成员保留历史，不能删除，可以停用。停用后不再出现在新的派发窗口中。"),
    ("改了超期天数，旧信件会怎样？",
     "保存后立即按新的天数重新判断。天数改小后，更多未办结信件会标为超期，并进入督办清单。"),
    FAQ_ALL[5],
    FAQ_ALL[6],
]


def faq_chapter(doc, items):
    doc.h1("五、常见问题")
    for i, (q, a) in enumerate(items, 1):
        doc.p(f"Q{i}：{q}")
        doc.p(f"A：{a}")


def groups_for(who):
    if who == "all":
        pc = [(PC_STU, "（学生）"), (PC_SEC, "（科室教师）"), (PC_ADM, "（综合办）")]
        mo = [(MO_STU, "（学生）"), (MO_SEC, "（科室教师）"), (MO_ADM, "（综合办）")]
    elif who == "stu":
        pc, mo = [(PC_STU, "")], [(MO_STU, "")]
    elif who == "sec":
        pc, mo = [(PC_SEC, "")], [(MO_SEC, "")]
    else:
        pc, mo = [(PC_ADM, "")], [(MO_ADM, "")]
    return pc, mo


def toc_lines(who):
    h12 = "1.2 使用角色" if who == "all" else "1.2 本角色要做的事"
    lines = [
        (1, "一、系统概述"),
        (2, "1.1 系统简介"),
        (2, h12),
        (2, "1.3 核心办理规则（重要）"),
        (1, "二、系统入口与登录"),
        (2, "2.1 登录方式"),
        (2, "2.2 PC 端入口"),
        (2, "2.3 移动端入口"),
        (2, "2.4 终端适用说明"),
        (1, "三、PC 端操作"),
    ]
    pc, mo = groups_for(who)
    n = 1
    for items, suffix in pc:
        for it in items:
            lines.append((2, f"3.{n} {it['title']}{suffix}"))
            n += 1
    lines.append((1, "四、移动端操作"))
    n = 1
    for items, suffix in mo:
        for it in items:
            lines.append((2, f"4.{n} {it['title']}{suffix}"))
            n += 1
    lines.append((1, "五、常见问题"))
    return lines


def build(filename, subtitle, who):
    doc = Manual()
    doc.cover(subtitle)
    doc.toc(toc_lines(who))
    overview(doc, who)
    entry_chapter(doc, who)
    pc, mo = groups_for(who)

    if who == "all":
        doc.h1("三、PC 端操作")
        doc.p("PC 端覆盖学生来信、科室办理和综合办受理配置。学生从“我要写信”开始，科室承办人从“科室工作台”办理派给本人的信件，综合办从“待分派信件”派发到具体承办老师。")
    elif who == "stu":
        doc.h1("三、PC 端操作")
        doc.p("学生在 PC 端完成写信、在我的信件跟踪进度、查看公开信件和评价。左侧菜单分为常用、信息公开和帮助。")
    elif who == "sec":
        doc.h1("三、PC 端操作")
        doc.p("科室承办人在左侧“本科室办理”中处理派发给本人的信件。公开信件和制度文件的查看方式与学生端相同。")
    else:
        doc.h1("三、PC 端操作")
        doc.p("综合办在 PC 端完成受理派发、督办和全部配置。菜单分为办理、运营、配置，以及本科室办理。")

    n = 1
    for items, suffix in pc:
        titled = [{**it, "title": it["title"] + suffix} for it in items]
        n = emit_steps(doc, 3, titled, n)

    if who == "all":
        doc.h1("四、移动端操作")
        doc.p("从智慧学工小程序进入学工信箱后，可办理日常来信、查看和单封派发。批量派发、批量催办和后台配置请使用 PC 端。")
    elif who == "stu":
        doc.h1("四、移动端操作")
        doc.p("移动端与 PC 端数据互通。进入学工信箱后，底部为“首页 / 公开 / 我的 / 更多”，写信前需阅读须知。")
    elif who == "sec":
        doc.h1("四、移动端操作")
        doc.p("进入学工信箱后办理派给本人的信件。底部为“工作台 / 收件箱 / 我的”。")
    else:
        doc.h1("四、移动端操作")
        doc.p("进入学工信箱后可查看待分派并单封派发。批量操作和配置请回到 PC 端。")

    n = 1
    for items, suffix in mo:
        titled = [{**it, "title": it["title"] + suffix} for it in items]
        n = emit_steps(doc, 4, titled, n)

    faq = {"all": FAQ_ALL, "stu": FAQ_STU, "sec": FAQ_SEC, "adm": FAQ_ADM}[who]
    faq_chapter(doc, faq)
    doc.page_numbers()
    doc.save(filename)


def main():
    build("学工信箱-操作手册（移动端+PC端）.docx",
          "操作手册（移动端 + PC 端）", "all")
    build("学工信箱-学生端操作手册（移动端+PC端）.docx",
          "学生端 · 操作手册（移动端 + PC 端）", "stu")
    build("学工信箱-科室教师端操作手册（移动端+PC端）.docx",
          "科室教师端 · 操作手册（移动端 + PC 端）", "sec")
    build("学工信箱-综合办管理员操作手册（移动端+PC端）.docx",
          "综合办管理员 · 操作手册（移动端 + PC 端）", "adm")


if __name__ == "__main__":
    main()
