#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《学工信箱修改清单》，同步最新角色、流程和消息。"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx_kit import Builder, set_run, FONT_HEAD
from mailbox_terms import (
    ROLE_NAMES, NOTE_KEZHANG, CHAIN, RECEIPT_NOTE, LIMIT_NOTE, FLOW_NOTE,
    COLLEGE_NOTE, APPROVAL_NOTE, STATUS_NOTE, REASSIGN_NOTE, WITHDRAW_NOTE,
    SECOND_NOTE, PUBLIC_1, PUBLIC_2, HANDLER, NODE_MAP, CANCELLED,
    CHANNEL_NOTE, PLACEHOLDERS, ACK, DONE, ROLES,
    FLOW as FLOW_ROWS, TEMPLATES, PLACEHOLDER_ROWS,
    assert_catalog, assert_docx, template_rows,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "学工信箱修改清单.docx")
FLOW_IMG = os.path.join(ROOT, "docs", "manual-assets", "mailbox-flow.png")


def lines(b, text):
    for line in text.split("\n"):
        b.body(line, indent=False)


def figure(b, path, caption):
    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(15.5))
    cap = b.doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    set_run(cap.add_run(caption), FONT_HEAD, 10.5, bold=True)


def main():
    assert_catalog()
    b = Builder()
    title = b.doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    set_run(title.add_run("学工信箱修改清单"), FONT_HEAD, 22, bold=True)
    b.note("2026年10月　·　供确认。角色为" + ROLE_NAMES + "。")
    b.body(NOTE_KEZHANG)
    b.body(CHAIN)

    b.h("一、办理口径", 1)
    b.table(
        ["序号", "事项", "确认口径"],
        [
            ["1", "回执", "学生提交后立即收到 7 日回执。回执不是正式答复，状态仍为待分派。草稿不进入办理。"],
            ["2", "分派", "综合办指定负责科室，把来信分派给该科室的科室负责人。不自动推荐科室，不直接指定科室回复人，也不直接分给科室分管领导。"],
            ["3", "科室负责人", "可以自己撰写；也可以指定科室回复人；也可以发起学院协同。后两者可以同时进行。"],
            ["4", "科室回复人", "只负责起草，写完交回科室负责人。不能自己提交审批，不能自己给学生发送。"],
            ["5", "审批", "回复稿回到科室负责人后，由负责人提交科室分管领导。不通过则退回负责人并附意见。负责人可自己改，或再交科室回复人改。分管领导不能改完直接发出。"],
            ["6", "回复学生", "审批通过后不自动发送。由科室负责人确认回复学生，并选择公开范围：全部公开，或仅该学院副书记可见。学生始终能看到自己的答复。"],
            ["7", "学院协同", "直接交给学生所在学院的副书记。副书记指定人员线下承办，写处理意见后发回科室负责人。不能对学生办结。线下承办不是系统审批。"],
            ["8", "二次办理", "办结后可评价。不满意则回到科室负责人，再审批，再由负责人确认发送。"],
            ["9", "学院人员", "学院副书记由学工系统维护。本板块不提供学院人员维护。"],
            ["10", "事务分类", "只用于写信和统计。不再维护科室承接事务参考、成员负责事务和分类科室职责参考。"],
        ],
        [1.4, 2.6, 12.0],
        font=10.5,
        aligns=["center", "center", "left"],
    )

    b.h("二、办理流程", 1)
    b.body(RECEIPT_NOTE)
    b.body(LIMIT_NOTE)
    b.body(FLOW_NOTE)
    b.body(COLLEGE_NOTE)
    b.body(APPROVAL_NOTE)
    b.body(STATUS_NOTE)
    figure(b, FLOW_IMG, "图 1  学工信箱办理流程")
    b.table(
        ["顺序", "办理人", "做什么", "学生所见", "办理环节"],
        FLOW_ROWS,
        [1.3, 2.6, 6.3, 2.4, 3.4],
        font=9,
        aligns=["center", "center", "left", "center", "left"],
    )
    b.body(REASSIGN_NOTE)
    b.body(WITHDRAW_NOTE)
    b.body(SECOND_NOTE)
    b.body(PUBLIC_1)
    b.body(PUBLIC_2)
    b.h("角色职责", 2)
    b.table(
        ["角色", "做什么", "不能做什么"],
        ROLES,
        [3.0, 6.8, 6.2],
        font=9,
        aligns=["center", "left", "left"],
    )

    b.h("三、消息提醒", 1)
    b.body("下面 %d 条为在用模板。标题与名称相同。回复稿交回科室负责人为新增编号 M-25，名称为「回复稿已返回」。" % len(TEMPLATES))
    b.body(NODE_MAP)
    b.body(CHANNEL_NOTE)
    b.body(HANDLER)
    b.body(CANCELLED)
    b.body(PLACEHOLDERS)
    b.table(
        ["编号", "名称（标题）", "发给谁", "何时发送", "正文要点"],
        template_rows(),
        [1.4, 2.3, 3.3, 3.6, 5.4],
        font=10,
        aligns=["center", "center", "left", "left", "left"],
    )
    b.h("占位符", 2)
    b.table(
        ["占位符", "发送时填入"],
        PLACEHOLDER_ROWS,
        [4.2, 11.8],
        font=10,
        aligns=["center", "left"],
    )
    b.h("学生可见正文", 2)
    b.body("下面两则是学生能看到的全文。花括号中的字是占位符，发送时按上表替换。M-01 不是正式答复。M-02 在科室负责人确认发送后才发出。")
    b.caption("M-01 提交回执")
    lines(b, ACK)
    b.caption("M-02 正式答复提醒")
    lines(b, DONE)

    b.h("四、功能上要同步的调整", 1)
    b.bullet("综合办改为指定负责科室，并分派给该科室的科室负责人。不自动推荐科室，不直接指定科室回复人，也不直接分给科室分管领导。")
    b.bullet("科室负责人可以自己写答复，也可以指定科室回复人，也可以发起学院协同。指定回复人和发起学院协同可以同时进行。")
    b.bullet("科室回复人只起草，写完把回复稿交回科室负责人。不能自己提交审批，不能自己给学生发送。")
    b.bullet("回复稿回到科室负责人后，由负责人提交科室分管领导审批。")
    b.bullet("审批不通过，退回科室负责人并附意见。负责人可自己改，或再交科室回复人改。分管领导不能改完直接发出。")
    b.bullet("审批通过后不自动发送。科室负责人确认回复学生，并选择全部公开，或仅该学院副书记可见。")
    b.bullet("学院协同直接发给学生所在学院的副书记。副书记指定人员线下承办，填写处理意见后发回科室负责人。不能对学生办结。")
    b.bullet("二次办理回到科室负责人，再审批，再由负责人确认发送。")
    b.bullet("不提供学院人员维护。事务分类只用于写信和统计。不再维护科室承接事务参考、成员负责事务和分类科室职责参考。")
    b.bullet("回执中的回复时限写为 7 日，与超期天数分开保存，当前取值都是 7 日。")
    b.bullet("不再单设信件查询。查看本人信件使用「我的信件」，查看已公开来信使用「公开信件」。")

    b.h("五、本次不改的内容", 1)
    b.bullet("登录入口不变：PC 端仍从中南大学官网「学生天地」进入智慧学工，移动端仍从「中南大学智慧学工」小程序进入。")
    b.bullet("学生写信字段不改。收悉回执不是满意度评价，也不是二次办理的起点。")

    b.save(OUT)
    doc = assert_docx(OUT, expect_toc=False)
    print("段落数: %d" % len(doc.paragraphs))
    print("表格数: %d" % len(doc.tables))
    print("模板条数: %d" % len(TEMPLATES))
    print("校验通过", OUT)


if __name__ == "__main__":
    main()
